"""매출 통계 이메일 빌더 및 Excel 생성"""
import os
from datetime import date
from email import encoders
from email.mime.base import MIMEBase

# ── 컬럼 정의 ────────────────────────────────────────────────────────────────
COLUMNS = [
    "구매경로", "매출", "수수료(순수익)", "주문건수",
    "누적회원수", "구매회원수", "구매회원비율",
    "기간내재구매회원수", "기간내재구매율",
    "누적재구매회원수", "누적재구매율",
]

# API 응답 키 → 컬럼명 매핑
_KEY_MAP = {
    "ordMdaNm":            "구매경로",
    "saleAmt":             "매출",
    "cms":                 "수수료(순수익)",
    "ordCnt":              "주문건수",
    "totalMbrCnt":         "누적회원수",
    "buyerCnt":            "구매회원수",
    "buyerRatio":          "구매회원비율",
    "repeatBuyerCnt":      "기간내재구매회원수",
    "repeatBuyerRatio":    "기간내재구매율",
    "cumRepeatBuyerCnt":   "누적재구매회원수",
    "cumRepeatBuyerRatio": "누적재구매율",
}

_FOOTNOTES = [
    "기간내재구매회원수: 해당 기간 안에서만 봤을 때, 같은 기간에 2회 이상 주문한 회원 수 (해당 기간에 국한된 반복구매)",
    "누적재구매회원수: 해당 기간 구매회원 중, 이전 기간 구매 이력까지 합산해서 지금까지 총 주문 횟수가 2회 이상인 회원 수 (과거 이력을 포함한 재구매자)",
]


def parse_sales_rows(data: list[dict]) -> list[dict]:
    """API data 배열 → 컬럼명 기반 dict 목록으로 변환"""
    rows = []
    for item in data:
        row = {}
        for api_key, col_name in _KEY_MAP.items():
            row[col_name] = item.get(api_key)
        rows.append(row)
    return rows


def _fmt_num(v) -> str:
    if v in (None, 0, ""):
        return "-"
    try:
        n = int(str(v).replace(",", ""))
        return f"{n:,}" if n != 0 else "-"
    except (ValueError, TypeError):
        return str(v)


def _fmt_pct(v) -> str:
    """소수 비율(0~1)을 퍼센트 문자열로"""
    if v in (None, "", 0):
        return "0.00%"
    try:
        return f"{float(v) * 100:.2f}%"
    except (ValueError, TypeError):
        return str(v)


def _fmt_cell(col: str, val) -> str:
    if col == "구매경로":
        return str(val) if val else "-"
    if col in ("구매회원비율", "기간내재구매율", "누적재구매율"):
        return _fmt_pct(val)
    return _fmt_num(val)


# ── HTML 빌더 ─────────────────────────────────────────────────────────────────

def _table_html(rows: list[dict], caption: str) -> str:
    th_cells = "".join(f"<th>{c}</th>" for c in COLUMNS)

    def row_html(r: dict, is_total: bool) -> str:
        style = " style='font-weight:bold;background:#fff9e6'" if is_total else ""
        cells = "".join(
            f"<td{style}>{_fmt_cell(c, r.get(c))}</td>"
            for c in COLUMNS
        )
        return f"<tr>{cells}</tr>"

    data_rows  = [r for r in rows if r.get("구매경로") != "합계"]
    total_rows = [r for r in rows if r.get("구매경로") == "합계"]
    body = "".join(row_html(r, False) for r in data_rows)
    body += "".join(row_html(r, True)  for r in total_rows)

    return f"""
<h2 style="font-size:16px;font-weight:bold;text-align:center;margin:28px 0 8px;color:#222">{caption}</h2>
<div style="overflow-x:auto">
<table style="width:100%;border-collapse:collapse;font-size:12px;min-width:960px">
  <thead>
    <tr style="background:#f5c400;color:#222;text-align:center">
      {th_cells}
    </tr>
  </thead>
  <tbody>{body}</tbody>
</table>
</div>"""


def build_sales_email_html(
    period_type: str,
    start: date,
    end: date,
    b2c_rows: list[dict],
    b2b_rows: list[dict],
) -> str:
    period_str = f"{start.strftime('%Y-%m-%d')} ~ {end.strftime('%Y-%m-%d')}"
    period_ko  = f"{start.strftime('%Y년 %m월 %d일')} ~ {end.strftime('%m월 %d일')}"
    label      = "주간" if period_type == "weekly" else "월간"
    period_tag = "주간" if period_type == "weekly" else "월간"

    b2c_table = _table_html(b2c_rows, f"매출 통계 - B2C (온누리몰) ({period_tag}, {period_str})")
    b2b_table = _table_html(b2b_rows, f"매출 통계 - B2B 통합 ({period_tag}, {period_str})")

    footnote_html = "".join(
        f"<p style='margin:4px 0;font-size:12px;color:#555'>- {ln}</p>"
        for ln in _FOOTNOTES
    )

    return f"""<!DOCTYPE html>
<html lang="ko">
<head><meta charset="UTF-8">
<style>
  body{{font-family:'Malgun Gothic',Arial,sans-serif;font-size:13px;color:#222;background:#fff;margin:0;padding:20px}}
  th,td{{padding:7px 10px;border:1px solid #ddd;text-align:right;white-space:nowrap}}
  th{{text-align:center;font-weight:bold}}
  td:first-child{{text-align:left}}
  tr:nth-child(even) td{{background:#fafafa}}
</style>
</head>
<body>
<p>안녕하세요. 캐스트이즈 도효빈입니다.</p>
<p>
  {period_ko} 온누리몰 {label} 매출 및 재구매 통계 데이터 입니다.<br>
  B2B몰 사이트별 상세 매출 현황은 첨부된 엑셀 파일로 확인 가능합니다.
</p>

{b2c_table}
{b2b_table}

<div style="margin-top:24px;padding:14px;background:#f9f9f9;border-left:3px solid #ccc;font-size:12px">
  {footnote_html}
</div>

<p style="margin-top:20px">감사합니다!<br>도효빈 드림</p>
</body>
</html>"""


# ── Excel 빌더 ────────────────────────────────────────────────────────────────

def build_b2b_excel(b2b_site_rows: list[dict], filepath: str) -> bool:
    """
    B2B 몰별 데이터를 Excel로 저장.
    b2b_site_rows: B2B_SITE gbCd 응답의 parsed rows (stNm 필드 포함)
    """
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        yellow = PatternFill("solid", fgColor="F5C400")
        bold   = Font(bold=True)
        center = Alignment(horizontal="center", vertical="center")
        right  = Alignment(horizontal="right")
        thin   = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # stNm 기준으로 그룹핑 (raw API item에 stNm 있음)
        sites = {}
        for row in b2b_site_rows:
            name = row.get("_stNm") or "B2B"
            sites.setdefault(name, []).append(row)

        for site_name, rows in sites.items():
            ws = wb.create_sheet(title=site_name[:31])

            ws.append(COLUMNS)
            for cell in ws[1]:
                cell.fill   = yellow
                cell.font   = bold
                cell.alignment = center
                cell.border = border

            for r in rows:
                row_vals = []
                for col in COLUMNS:
                    val = r.get(col)
                    if col in ("구매회원비율", "기간내재구매율", "누적재구매율"):
                        try:
                            val = round(float(val) * 100, 2) if val else 0
                        except (TypeError, ValueError):
                            val = 0
                    elif col not in ("구매경로",):
                        try:
                            val = int(val) if val else 0
                        except (TypeError, ValueError):
                            pass
                    row_vals.append(val)
                ws.append(row_vals)
                is_total = r.get("구매경로") == "합계"
                if is_total:
                    for cell in ws[ws.max_row]:
                        cell.font = bold

            # 열 너비 자동 조정
            for col_cells in ws.columns:
                length = max((len(str(c.value or "")) for c in col_cells), default=8) + 2
                ws.column_dimensions[col_cells[0].column_letter].width = min(length, 22)

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.border = border
                    if cell.column > 1:
                        cell.alignment = right

        wb.save(filepath)
        return True
    except Exception as e:
        print(f"[Excel 생성 실패] {e}")
        return False


def attach_excel(msg, filepath: str, filename: str) -> None:
    with open(filepath, "rb") as f:
        part = MIMEBase(
            "application",
            "vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", "attachment", filename=filename)
    msg.attach(part)
